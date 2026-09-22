"""Load documents and extract their text, page by page."""

from dataclasses import dataclass
from pathlib import Path
from typing import Callable

from docx import Document as DocxDocument
from pypdf import PdfReader


class DocumentLoadError(Exception):
    """Raised when a document cannot be read."""


class UnsupportedFileTypeError(DocumentLoadError):
    """Raised when the file type is not supported."""


@dataclass
class PageText:
    """Text from one page. page_number is None for formats without pages."""

    page_number: int | None
    text: str


def _load_pdf(path: Path) -> list[PageText]:
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        raise DocumentLoadError(f"{path.name} is password-protected.")
    return [
        PageText(page_number=i, text=page.extract_text() or "")
        for i, page in enumerate(reader.pages, start=1)
    ]


def _load_docx(path: Path) -> list[PageText]:
    doc = DocxDocument(str(path))
    parts = [p.text for p in doc.paragraphs]
    # Include text inside tables too
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return [PageText(page_number=None, text="\n".join(parts))]


def _load_text(path: Path) -> list[PageText]:
    text = path.read_text(encoding="utf-8", errors="replace")
    return [PageText(page_number=None, text=text)]


# Map each file extension to its loader. Add new types here later.
LOADERS: dict[str, Callable[[Path], list[PageText]]] = {
    ".pdf": _load_pdf,
    ".docx": _load_docx,
    ".txt": _load_text,
    ".md": _load_text,
}

SUPPORTED_EXTENSIONS = set(LOADERS)


def load_document(path: Path) -> list[PageText]:
    """Read a document and return its text, one entry per page."""
    extension = path.suffix.lower()
    if extension not in LOADERS:
        raise UnsupportedFileTypeError(
            f"'{extension}' files are not supported. "
            f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
        )
    if not path.exists():
        raise DocumentLoadError(f"File not found: {path.name}")

    try:
        pages = LOADERS[extension](path)
    except DocumentLoadError:
        raise
    except Exception as error:  # corrupted or unreadable file
        raise DocumentLoadError(f"Could not read {path.name}: {error}") from error

    if not any(page.text.strip() for page in pages):
        raise DocumentLoadError(
            f"No text found in {path.name}. If it is a scanned PDF, "
            "it contains images of text, which is not supported yet."
        )
    return pages